import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Преобразует CSV-файл в XLSX.
    Первая строка CSV — заголовок.
    Лист называется Sheet1.
    Ширина колонок подстраивается под текст.
    """

    if not csv_path.endswith(".csv"):
        raise ValueError("Файл должен быть CSV")
    if not xlsx_path.endswith(".xlsx"):
        raise ValueError("Файл должен быть XLSX")

    with open(csv_path, "r", encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)

    if len(rows) < 2:
        raise ValueError("CSV пустой или содержит только заголовок")
    if not rows[0] or not any(rows[0]):
        raise ValueError("CSV без заголовка")

    for name in rows[0]:
        if not name:
            raise ValueError("Пустое название столбца")
    for row in rows:
        if len(row) != len(rows[0]):
            raise ValueError("Неверное количество столбцов CSV")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)
    for i in range(len(rows[0])):
        max_length = 8

        for row in rows:
            if len(row[i]) > max_length:
                max_length = len(row[i])
        letter = get_column_letter(i + 1)
        ws.column_dimensions[letter].width = max_length + 2

    path = Path(xlsx_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb.save(path)

if __name__ == "__main__":
    csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")